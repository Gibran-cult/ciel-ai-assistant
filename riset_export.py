"""Ekspor hasil Riset Skripsi (peta sumber) ke DOCX.

Flow "Ciel Riset Skripsi" selalu menjawab dengan format tetap: judul `#`,
ringkasan, bagian `##` (tema, celah riset, pertanyaan untuk dosen), lalu
"Daftar Pustaka" yang disusun komponen flow dari metadata Exa. Modul ini
hanya mengenal subset markdown itu: heading, bullet, daftar bernomor,
paragraf, **tebal**, *miring*, dan tautan [teks](url).

Format dokumen mengikuti kebiasaan skripsi di Indonesia: Times New Roman
12, spasi 1,5, dan daftar pustaka dengan indentasi gantung.
"""

import io
import re

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
BODY_PT = 12
HEADING_PT = {1: 14, 2: 12, 3: 12}
HANGING_INDENT = Cm(1.27)

HEADING = re.compile(r"^(#{1,3})\s+(.*)$")
BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")
NUMBERED = re.compile(r"^\s*(\d+)[.)]\s+(.*)$")
REFERENCES_HEADING = re.compile(r"^daftar\s+pustaka\b", re.IGNORECASE)
# **tebal**, *miring*, [teks](url); sisanya teks biasa.
INLINE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|\[([^\]]+)\]\((https?://[^)\s]+)\)")

FOOTNOTE = (
    "Disusun dengan Ciel Riset Skripsi. Peta sumber ini adalah titik awal, "
    "bukan isi skripsi: buka dan baca setiap sumber sebelum mengutipnya."
)


def inline_runs(text):
    """Teks satu baris -> [(teks, tebal, miring)]. Tautan ditulis "teks (url)"."""

    runs = []
    position = 0

    for match in INLINE.finditer(text):
        if match.start() > position:
            runs.append((text[position:match.start()], False, False))

        bold, italic, label, url = match.groups()

        if bold is not None:
            runs.append((bold, True, False))
        elif italic is not None:
            runs.append((italic, False, True))
        elif label == url:
            runs.append((url, False, False))
        else:
            runs.append((f"{label} ({url})", False, False))

        position = match.end()

    if position < len(text):
        runs.append((text[position:], False, False))

    return runs


def parse_blocks(markdown):
    """
    Markdown -> [(jenis, teks)] dengan jenis h1/h2/h3, bullet, number,
    paragraph, atau reference (baris di bawah heading Daftar Pustaka).
    Baris yang bersebelahan dalam satu paragraf digabung.
    """

    blocks = []
    in_references = False
    paragraph = []

    def flush():
        if paragraph:
            blocks.append(("paragraph", " ".join(paragraph)))
            paragraph.clear()

    for raw in str(markdown).splitlines():
        line = raw.strip()

        if not line:
            flush()
            continue

        heading = HEADING.match(line)

        if heading:
            flush()
            text = heading.group(2).strip()
            in_references = bool(REFERENCES_HEADING.match(text))
            blocks.append((f"h{len(heading.group(1))}", text))
            continue

        bullet = BULLET.match(line)
        numbered = NUMBERED.match(line)

        entry = bullet.group(1) if bullet else line

        if in_references and entry.startswith("["):
            # Satu entri pustaka per baris ("[n] ..."), termasuk yang ditulis sebagai bullet.
            flush()
            blocks.append(("reference", entry))
        elif bullet:
            flush()
            blocks.append(("bullet", bullet.group(1)))
        elif numbered:
            flush()
            blocks.append(("number", f"{numbered.group(1)}. {numbered.group(2)}"))
        else:
            paragraph.append(line)

    flush()
    return blocks


def _style_document(document):
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)
    # Nama font Asia Timur ikut diatur agar Word tidak menggantinya.
    normal.element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", FONT
    )

    for level, size in HEADING_PT.items():
        style = document.styles[f"Heading {level}"]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def _write_runs(paragraph, text):
    for chunk, bold, italic in inline_runs(text):
        run = paragraph.add_run(chunk)
        run.bold = bold or None
        run.italic = italic or None


def report_docx(markdown, topic):
    """Hasil Riset Skripsi (markdown) -> isi file .docx (bytes)."""

    document = Document()
    _style_document(document)
    document.core_properties.title = f"Peta Sumber: {topic}"[:255]
    document.core_properties.author = "Ciel Riset Skripsi"

    for kind, text in parse_blocks(markdown):

        if kind.startswith("h"):
            document.add_heading(text, level=int(kind[1]))
            continue

        style = {"bullet": "List Bullet"}.get(kind)
        paragraph = document.add_paragraph(style=style)
        _write_runs(paragraph, text)

        if kind == "reference":
            paragraph.paragraph_format.left_indent = HANGING_INDENT
            paragraph.paragraph_format.first_line_indent = -HANGING_INDENT

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    run = note.add_run(FOOTNOTE)
    run.italic = True
    run.font.size = Pt(10)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def report_filename(topic):
    """Nama file aman untuk unduhan: "Peta Sumber - Literasi Digital UMKM.docx"."""
    name = ""

    # Dipotong per kata (maks. 60 karakter), bukan di tengah kata.
    for word in re.sub(r"[^\w\s-]", " ", str(topic)).split():
        if name and len(name) + 1 + len(word) > 60:
            break
        name = f"{name} {word}".strip()[:60]

    name = name or "Riset Skripsi"
    return f"Peta Sumber - {name}.docx"
