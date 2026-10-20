"""writing.py: Kerangka Bab 1-3 dari peta sumber."""

import io
import sys
import unittest
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import riset_export  # noqa: E402
import writing  # noqa: E402

SOURCE_MAP = """# Peta Sumber: Literasi digital UMKM

Ringkasan [1][2].

## Tema: Adopsi teknologi
- Temuan [1].

## Daftar Pustaka

[1] Sari, D. (2023). *Literasi Digital*. Jurnal X. https://example.org/a

[2] Putra, A. (2024). *Adopsi E-commerce*. Jurnal Y. https://example.org/b

[3] Kemenkop. (t.t.). *Data UMKM*. kemenkop.go.id. https://kemenkop.go.id

Daftar pustaka disusun otomatis dari data sumber hasil pencarian."""

ANSWER = """# Kerangka Bab 1–3: Literasi digital UMKM

## BAB I PENDAHULUAN
### 1.1 Latar Belakang
- Literasi digital naikkan kinerja [1][7]
- Data dukungan pemerintah [3, 9]
- Klaim tanpa sumber yang dikenal [8].

## Daftar Pustaka
[1] Sumber karangan model. (2099). https://karangan.example
"""


class KerangkaTest(unittest.TestCase):

    def test_allowed_numbers_and_reference_section_come_from_the_map(self):
        self.assertEqual(writing.allowed_numbers(SOURCE_MAP), {"1", "2", "3"})
        self.assertTrue(writing.reference_section(SOURCE_MAP).startswith("## Daftar Pustaka"))
        self.assertEqual(writing.reference_section("tanpa daftar"), "")
        self.assertEqual(writing.allowed_numbers("tanpa daftar"), set())
        self.assertEqual(writing.map_topic(SOURCE_MAP), "Literasi digital UMKM")

    def test_payload_marks_the_mode_and_carries_the_whole_map(self):
        payload = writing.kerangka_payload(SOURCE_MAP, "Literasi digital UMKM", "Kuantitatif", "UMKM Bandung")
        self.assertTrue(payload.startswith("[KERANGKA]\nTopik skripsi: Literasi digital UMKM\nPendekatan penelitian: Kuantitatif"))
        self.assertIn("Objek atau lokasi penelitian: UMKM Bandung", payload)
        self.assertIn("PETA SUMBER:\n# Peta Sumber: Literasi digital UMKM", payload)
        self.assertIn("[3] Kemenkop.", payload)
        self.assertNotIn("Objek atau lokasi", writing.kerangka_payload(SOURCE_MAP, "x", "Belum tahu"))

    def test_unknown_citations_are_dropped(self):
        allowed = {"1", "2", "3"}
        self.assertEqual(writing.keep_known_citations("A [1][7]", allowed), "A [1]")
        self.assertEqual(writing.keep_known_citations("B [3, 9]", allowed), "B [3]")
        self.assertEqual(writing.keep_known_citations("C tanpa sumber [8].", allowed), "C tanpa sumber.")
        self.assertEqual(writing.keep_known_citations("D [2, 3]", allowed), "D [2, 3]")

    def test_finished_kerangka_uses_only_the_maps_sources(self):
        text = writing.finish_kerangka(ANSWER, SOURCE_MAP)
        body, _, refs = text.partition("## Daftar Pustaka (dari peta sumber)")
        # The model's own reference list is gone; the map's list is appended.
        self.assertNotIn("karangan", text)
        self.assertNotIn("[7]", body)
        self.assertNotIn("[9]", body)
        self.assertNotIn("[8]", body)
        self.assertIn("- Literasi digital naikkan kinerja [1]", body)
        self.assertIn("- Data dukungan pemerintah [3]", body)
        self.assertIn("[2] Putra, A. (2024).", refs)
        self.assertEqual(text.count("## Daftar Pustaka"), 1)

    def test_kerangka_docx_has_its_own_label_footnote_and_name(self):
        text = writing.finish_kerangka(ANSWER, SOURCE_MAP)
        data = riset_export.report_docx(text, "Literasi digital UMKM", label="Kerangka Bab 1–3",
                                        footnote=writing.KERANGKA_FOOTNOTE)
        document = docx.Document(io.BytesIO(data))
        self.assertEqual(document.core_properties.title, "Kerangka Bab 1–3: Literasi digital UMKM")
        self.assertIn("bukan teks skripsi", document.paragraphs[-1].text)
        references = [p for p in document.paragraphs if p.text.startswith("[")]
        self.assertEqual(len(references), 3)
        self.assertEqual(riset_export.report_filename("Literasi digital UMKM", prefix="Kerangka Bab 1-3"),
                         "Kerangka Bab 1-3 - Literasi digital UMKM.docx")
        # The peta sumber keeps its old title and name.
        self.assertEqual(riset_export.report_filename("X"), "Peta Sumber - X.docx")


if __name__ == "__main__":
    unittest.main()
