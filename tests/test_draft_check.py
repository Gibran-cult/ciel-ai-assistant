import base64
import io
import json
import unittest

from docx import Document

from draft_check import (
    MAX_DRAFT_CHARS,
    check_citations,
    check_report,
    clean_draft,
    draft_from_file,
    find_citations,
    parse_references,
    similar_sentences,
    split_excerpts,
    without_own_references,
)

SOURCE_MAP = """# Peta Sumber: Literasi digital UMKM

Ringkasan singkat.

## Tema: Literasi digital dan kinerja

- Literasi digital berpengaruh positif signifikan terhadap kinerja UMKM kuliner di Kota Padang menurut regresi linier berganda [1].
- Orientasi kewirausahaan memperkuat hubungan literasi digital dengan kinerja usaha [2].

## Daftar Pustaka

[1] Nurlina Nurlina, Yenni Del Rosa, Budi Yanti. (2023). *The Effect Of Digital Literacy*. Jurnal X. https://example.org/1

[2] Sinta Fajri Utami, Fitra Dila Lestari. (2026). *Pengaruh Literasi Digital*. doi.org. https://doi.org/10.1/abc

[3] *Statistik UMKM Indonesia*. (t.t.). bps.go.id. https://bps.go.id/umkm

Daftar pustaka disusun otomatis dari data sumber hasil pencarian.
"""

EXCERPT = (
    "Hasil penelitian menunjukkan bahwa literasi digital memiliki pengaruh yang kuat "
    "terhadap kemampuan pelaku usaha mikro dalam memasarkan produk secara daring."
)


def marker(items):
    encoded = base64.urlsafe_b64encode(json.dumps(items).encode("utf-8")).decode("ascii")
    return f"\n\n<!-- ciel-excerpts:v1 {encoded} -->\n"


class SplitExcerptsTest(unittest.TestCase):
    def test_marker_is_removed_and_parsed(self):
        text, excerpts = split_excerpts(SOURCE_MAP + marker([{"n": 1, "highlights": [EXCERPT, " "]}, {"n": 2, "highlights": []}]))
        self.assertEqual(text, SOURCE_MAP.rstrip())
        self.assertEqual(excerpts, {"1": [EXCERPT]})

    def test_answer_without_marker_is_unchanged(self):
        self.assertEqual(split_excerpts(SOURCE_MAP), (SOURCE_MAP, {}))

    def test_broken_marker_is_removed_without_excerpts(self):
        text, excerpts = split_excerpts(SOURCE_MAP + "\n<!-- ciel-excerpts:v1 bm90LWpzb24= -->")
        self.assertEqual((text, excerpts), (SOURCE_MAP.rstrip(), {}))

    def test_marker_in_the_middle_is_not_trusted(self):
        answer = "Teks" + marker([{"n": 1, "highlights": ["x"]}]) + "lanjutan"
        self.assertEqual(split_excerpts(answer), (answer, {}))


class ReferencesTest(unittest.TestCase):
    def test_parse_references_from_the_map(self):
        references = parse_references(SOURCE_MAP)
        self.assertEqual(set(references), {"1", "2", "3"})
        self.assertEqual(references["1"]["year"], "2023")
        self.assertEqual(references["1"]["label"], "Nurlina dkk., 2023")
        self.assertEqual(references["3"]["authors"], "")
        self.assertEqual(references["3"]["year"], "t.t.")
        self.assertEqual(references["3"]["title"], "Statistik UMKM Indonesia")


class CitationsTest(unittest.TestCase):
    def setUp(self):
        self.references = parse_references(SOURCE_MAP)

    def test_finds_numbered_parenthetical_and_narrative_citations(self):
        draft = "Satu [1, 3]. Dua (Nurlina dkk., 2023; Utami, 2026). Menurut Utami dkk. (2026), tiga."
        texts = [text for text, _, _ in find_citations(draft)]
        self.assertEqual(texts, ["[1]", "[3]", "(Nurlina dkk., 2023)", "(Utami, 2026)", "Utami dkk. (2026)"])

    def test_number_ranges_are_expanded(self):
        self.assertEqual([v for _, _, v in find_citations("lihat [1–3]")], ["1", "2", "3"])

    def test_plain_years_in_parentheses_are_not_citations(self):
        self.assertEqual(find_citations("Data (sejak 2020 hingga 2024) dan tahun (2023) naik."), [])

    def test_citations_are_matched_by_name_and_year(self):
        draft = (
            "A (Nurlina, 2023). B (Utami & Lestari, 2025). C (Santoso, 2022). "
            "D [9]. E [2]. F Nurlina et al. (2023)."
        )
        result = check_citations(draft, self.references)
        self.assertEqual(result["ok"], {"1": ["(Nurlina, 2023)", "Nurlina et al. (2023)"], "2": ["[2]"]})
        self.assertEqual(result["year"], [("(Utami & Lestari, 2025)", "2")])
        self.assertEqual(result["missing"], ["(Santoso, 2022)"])
        self.assertEqual(result["unknown_numbers"], ["[9]"])

    def test_accents_and_case_do_not_matter(self):
        references = {"1": {"authors": "José Pérez", "title": "", "year": "2024", "label": "Pérez, 2024"}}
        self.assertEqual(check_citations("x (PEREZ, 2024).", references)["ok"], {"1": ["(PEREZ, 2024)"]})


class SimilarityTest(unittest.TestCase):
    def test_copied_excerpt_sentence_is_flagged(self):
        draft = EXCERPT + " Kalimat ini adalah tulisan sendiri yang tidak meniru sumber mana pun sama sekali."
        flagged = similar_sentences(draft, SOURCE_MAP, {"4": [EXCERPT]})
        self.assertEqual(len(flagged), 1)
        self.assertEqual(flagged[0][1], 1.0)
        self.assertEqual(flagged[0][2], "cuplikan sumber [4]")

    def test_copied_map_summary_is_flagged_without_its_marker(self):
        draft = "Literasi digital berpengaruh positif signifikan terhadap kinerja UMKM kuliner di Kota Padang menurut regresi linier berganda."
        flagged = similar_sentences(draft, SOURCE_MAP, {})
        self.assertEqual([label for _, _, label in flagged], ["ringkasan peta sumber Ciel"])

    def test_words_of_the_topic_do_not_count_as_copying(self):
        # Seen on the live app: a sentence restating the thesis title scored 44%
        # against the map, whose heading and summary repeat the same title.
        source_map = (
            "# Peta Sumber: Pengaruh Media Sosial terhadap Minat Beli Mahasiswa\n\n"
            "- Affiliate marketing berpengaruh positif terhadap minat beli mahasiswa di TikTok Shop [6]\n"
        )
        draft = "Penelitian ini menguji pengaruh media sosial terhadap minat beli mahasiswa di kampus kami."
        self.assertEqual(len(similar_sentences(draft, source_map, {})), 1)
        self.assertEqual(similar_sentences(draft, source_map, {}, "Pengaruh media sosial terhadap minat beli mahasiswa"), [])

    def test_direct_quote_and_short_sentences_are_not_flagged(self):
        draft = f"“{EXCERPT}” [1]. Literasi digital berpengaruh."
        self.assertEqual(similar_sentences(draft, SOURCE_MAP, {"1": [EXCERPT]}), [])

    def test_reference_list_of_the_draft_is_ignored(self):
        draft = "Isi bab.\n\nDaftar Pustaka\n[1] Nurlina Nurlina. (2023). The Effect Of Digital Literacy."
        self.assertEqual(without_own_references(draft), "Isi bab.\n\n")


class DraftInputTest(unittest.TestCase):
    def test_docx_and_text_files_are_read(self):
        document = Document()
        document.add_paragraph("Paragraf satu.")
        document.add_paragraph("Paragraf dua.")
        raw = io.BytesIO()
        document.save(raw)
        self.assertEqual(draft_from_file("Bab 1.DOCX", raw.getvalue()), "Paragraf satu.\nParagraf dua.")
        self.assertEqual(draft_from_file("bab.txt", "Halo".encode("utf-8")), "Halo")

    def test_broken_files_raise_value_error(self):
        with self.assertRaises(ValueError):
            draft_from_file("bab.docx", b"bukan docx")
        with self.assertRaises(ValueError):
            draft_from_file("bab.txt", b"\x00\x01")

    def test_clean_draft_collapses_spaces_and_caps_length(self):
        self.assertEqual(clean_draft("  a   b \n\n c "), "a b\n\nc")
        self.assertEqual(len(clean_draft("x" * (MAX_DRAFT_CHARS + 10))), MAX_DRAFT_CHARS)


class ReportTest(unittest.TestCase):
    def test_report_lists_problems_matches_and_copied_sentences(self):
        draft = f"{EXCERPT} (Utami, 2025). Penelitian ini dimulai dari tulisan sendiri [1] dan (Santoso, 2022)."
        report = check_report(draft, SOURCE_MAP, {"2": [EXCERPT]}, "Literasi *digital*")
        self.assertTrue(report.startswith("# Cek Kutipan & Kemiripan: Literasi \\*digital\\*"))
        self.assertIn("cuplikan 1 sumbernya", report)
        self.assertIn("Sitasi yang cocok dengan peta sumber: 1. Perlu dicek: 2.", report)
        self.assertIn("namanya cocok dengan [2] (Utami dkk., 2026), tapi tahunnya beda", report)
        self.assertIn("(Santoso, 2022): tidak ada di peta sumber", report)
        self.assertIn("- [1] Nurlina dkk., 2023: dikutip 1 kali.", report)
        self.assertIn("100% rangkaian katanya sama dengan cuplikan sumber [2]", report)
        self.assertIn("Sumber peta yang belum kamu kutip: [3].", report)
        self.assertIn("bukan dengan seluruh internet seperti Turnitin", report)

    def test_clean_draft_says_so_and_old_maps_are_explained(self):
        draft = "Penelitian ini membahas topik yang menarik dengan kata-kata sendiri sepenuhnya tanpa salinan apa pun [1]."
        report = check_report(draft, SOURCE_MAP, {}, "Topik")
        self.assertIn("sebelum Ciel menyimpan cuplikan sumber", report)
        self.assertIn("Tidak ada sitasi yang janggal atau kalimat yang terlalu mirip", report)
        self.assertIn("- Kalimat yang terlalu mirip: tidak ada.", report)


if __name__ == "__main__":
    unittest.main()
